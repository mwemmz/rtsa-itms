"""Developer 2 platform tests: security, RBAC, payments, notifications,
reports/exports, inter-agency integration, operations."""

import gzip
import hashlib
import hmac
import io
import json
from datetime import datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

import app.core.ratelimit as ratelimit
from app.core import totp
from app.core.database import SessionLocal
from app.models.enforcement import Challan, ChallanStatus, Violation, ViolationType
from app.models.notification import NotificationRule
from app.models.user import User
from app.models.vehicle import Vehicle
from app.services import integration as integration_service
from app.services import settings as runtime_settings
from main import app
from tests.conftest import create_user

client = TestClient(app)
PW = "password123"


@pytest.fixture(autouse=True)
def _reset_state():
    yield
    with ratelimit._lock:
        ratelimit._attempts.clear()
    integration_service.reset_throttle()
    runtime_settings.invalidate_cache()


def _login(email, password=PW, **headers):
    r = client.post("/api/auth/login", json={"email": email, "password": password}, headers=headers)
    return r


def _auth_headers(role="admin", **extra):
    email, uid = create_user(role, **extra)
    r = _login(email)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}, email, uid


def _db():
    return SessionLocal()


def _rule(event, channels="in_app", title="t", body="b {reference}"):
    db = _db()
    try:
        if not db.query(NotificationRule).filter(NotificationRule.trigger_event == event).first():
            db.add(NotificationRule(trigger_event=event, channels=channels, title_template=title, body_template=body))
            db.commit()
    finally:
        db.close()


def _vehicle_with_fine(owner_id, amount=50000, plate=None):
    db = _db()
    try:
        v = Vehicle(registration_number=plate or f"T{uuid4().hex[:6].upper()}", owner_name="Own Er",
                    owner_id_number="1", make="Toyota", model="Hilux", year=2020, user_id=owner_id)
        db.add(v)
        db.flush()
        viol = Violation(vehicle_id=v.id, violation_type=ViolationType.SPEEDING, location="Great East Rd")
        db.add(viol)
        db.flush()
        c = Challan(reference=f"CH-{uuid4().hex[:8].upper()}", violation_id=viol.id, vehicle_id=v.id,
                    penalty_amount=amount, due_date=datetime.utcnow() + timedelta(days=10))
        db.add(c)
        db.commit()
        return str(v.id), str(c.id)
    finally:
        db.close()


# ============================ security ============================

def test_public_registration_cannot_create_staff():
    r = client.post("/api/auth/register", json={"email": f"x{uuid4().hex[:6]}@t.com", "password": PW,
                                                "full_name": "Sneaky", "role": "admin"})
    assert r.status_code == 403
    r = client.post("/api/auth/register", json={"email": f"x{uuid4().hex[:6]}@t.com", "password": PW,
                                                "full_name": "Ok User"})
    assert r.status_code == 201 and r.json()["role"] == "citizen"


def test_password_policy_on_register():
    r = client.post("/api/auth/register", json={"email": f"x{uuid4().hex[:6]}@t.com", "password": "short1",
                                                "full_name": "A"})
    assert r.status_code == 400
    r = client.post("/api/auth/register", json={"email": f"x{uuid4().hex[:6]}@t.com", "password": "onlyletters",
                                                "full_name": "A"})
    assert r.status_code == 400


def test_account_lockout_after_repeated_failures_and_admin_unlock():
    admin, _, _ = _auth_headers("admin")
    email, uid = create_user("citizen")
    for _ in range(5):
        # vary the client IP so only the per-account lockout (not the IP throttle) is exercised
        r = _login(email, "wrong-password1", **{"X-Forwarded-For": f"10.0.0.{_}"})
        assert r.status_code == 401
    r = _login(email, PW, **{"X-Forwarded-For": "10.9.9.9"})
    assert r.status_code == 423  # locked even with the right password
    r = client.post(f"/api/admin/users/{uid}/unlock", headers=admin)
    assert r.status_code == 200
    assert _login(email, PW, **{"X-Forwarded-For": "10.9.9.8"}).status_code == 200


def test_login_attempts_are_recorded():
    admin, _, _ = _auth_headers("admin")
    email, _ = create_user("citizen")
    _login(email, "nope-nope1")
    _login(email)
    r = client.get(f"/api/admin/login-attempts?email={email}", headers=admin)
    assert r.status_code == 200
    outcomes = [a["success"] for a in r.json()]
    assert True in outcomes and False in outcomes


def test_logout_revokes_the_session():
    h, _, _ = _auth_headers("citizen")
    assert client.get("/api/auth/me", headers=h).status_code == 200
    assert client.post("/api/auth/logout", headers=h).status_code == 204
    assert client.get("/api/auth/me", headers=h).status_code == 401


def test_token_without_session_is_rejected():
    from app.core.security import create_access_token

    _, uid = create_user("citizen")
    forged = create_access_token({"sub": str(uid), "role": "admin"})  # no sid
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {forged}"}).status_code == 401


def test_idle_session_times_out():
    h, _, _ = _auth_headers("citizen")
    from app.models.platform import UserSession

    db = _db()
    try:
        for s in db.query(UserSession).filter(UserSession.revoked_at.is_(None)):
            s.last_seen_at = datetime.utcnow() - timedelta(hours=3)
        db.commit()
    finally:
        db.close()
    r = client.get("/api/auth/me", headers=h)
    assert r.status_code == 401 and "inactivity" in r.json()["detail"]


def test_session_list_and_revoke_others():
    email, _ = create_user("citizen")
    t1 = _login(email).json()["access_token"]
    t2 = _login(email, **{"User-Agent": "Mozilla/5.0 (Android) Chrome/120"}).json()["access_token"]
    h1 = {"Authorization": f"Bearer {t1}"}
    sessions = client.get("/api/auth/sessions", headers=h1).json()
    assert len(sessions) >= 2 and sum(s["is_current"] for s in sessions) == 1
    assert client.post("/api/auth/sessions/revoke-others", headers=h1).json()["revoked"] >= 1
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {t2}"}).status_code == 401
    assert client.get("/api/auth/me", headers=h1).status_code == 200


def test_device_tracking_and_blocking():
    email, _ = create_user("citizen")
    _login(email, **{"User-Agent": "Mozilla/5.0 (Windows) Firefox/120"})
    t = _login(email, **{"User-Agent": "Mozilla/5.0 (Android) Chrome/120"}).json()["access_token"]
    h = {"Authorization": f"Bearer {t}"}
    devices = client.get("/api/auth/devices", headers=h).json()
    assert len(devices) == 2
    other = next(d for d in devices if "Firefox" in d["label"])
    r = client.patch(f"/api/auth/devices/{other['id']}?blocked=true", headers=h)
    assert r.status_code == 200 and r.json()["is_blocked"]
    # the blocked device can no longer sign in
    assert _login(email, **{"User-Agent": "Mozilla/5.0 (Windows) Firefox/120"}).status_code == 403


def test_mfa_enrolment_login_and_recovery_code():
    h, email, _ = _auth_headers("officer")
    setup = client.post("/api/auth/mfa/setup", headers=h).json()
    secret = setup["secret"]
    assert setup["otpauth_uri"].startswith("otpauth://totp/")
    assert client.post("/api/auth/mfa/enable", json={"code": "000000"}, headers=h).status_code == 400
    enabled = client.post("/api/auth/mfa/enable", json={"code": totp.totp_now(secret)}, headers=h)
    assert enabled.status_code == 200
    recovery = enabled.json()["recovery_codes"]
    assert len(recovery) == 8

    # password alone no longer yields a session
    r = _login(email)
    body = r.json()
    assert body["mfa_required"] is True and body["access_token"] is None
    bad = client.post("/api/auth/mfa/verify", json={"mfa_token": body["mfa_token"], "code": "123456"})
    assert bad.status_code == 401
    good = client.post("/api/auth/mfa/verify", json={"mfa_token": body["mfa_token"], "code": totp.totp_now(secret)})
    assert good.status_code == 200 and good.json()["access_token"]
    # an MFA challenge token is not an access token
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {body['mfa_token']}"}).status_code == 401

    # a recovery code works once
    body = _login(email).json()
    ok = client.post("/api/auth/mfa/verify", json={"mfa_token": body["mfa_token"], "code": recovery[0]})
    assert ok.status_code == 200
    body = _login(email).json()
    again = client.post("/api/auth/mfa/verify", json={"mfa_token": body["mfa_token"], "code": recovery[0]})
    assert again.status_code == 401


def test_admin_must_enrol_own_mfa_before_requiring_it_platform_wide():
    admin, _, _ = _auth_headers("admin")
    # refuses to flip the switch while the acting admin - a staff role too - has no MFA,
    # which would otherwise lock them out of Settings the moment it takes effect
    r = client.put("/api/admin/settings/security.require_mfa_staff", json={"value": True}, headers=admin)
    assert r.status_code == 400 and "mfa" in r.json()["detail"].lower()


def test_hard_mfa_enforcement_blocks_staff_until_enrolled():
    admin, _, admin_id = _auth_headers("admin")
    setup = client.post("/api/auth/mfa/setup", headers=admin).json()
    assert client.post("/api/auth/mfa/enable", json={"code": totp.totp_now(setup["secret"])},
                        headers=admin).status_code == 200
    assert client.put("/api/admin/settings/security.require_mfa_staff", json={"value": True},
                       headers=admin).status_code == 200
    try:
        h, email, _ = _auth_headers("officer")
        # logging in succeeds (no MFA enabled yet, so no mfa_required challenge)...
        r = _login(email)
        assert r.status_code == 200 and r.json()["access_token"]
        assert r.json()["mfa_setup_required"] is True
        # ...but every endpoint except the MFA-setup ones is blocked
        assert client.get("/api/notifications/", headers=h).status_code == 403
        assert client.get("/api/notifications/", headers=h).json()["detail"].lower().count("mfa") > 0
        assert client.get("/api/auth/me", headers=h).json()["mfa_setup_required"] is True
        # citizens are unaffected by the staff-only policy
        ch, cemail, _ = _auth_headers("citizen")
        assert client.get("/api/notifications/", headers=ch).status_code == 200
        # enrol, then the officer regains full access
        setup = client.post("/api/auth/mfa/setup", headers=h).json()
        enabled = client.post("/api/auth/mfa/enable", json={"code": totp.totp_now(setup["secret"])}, headers=h)
        assert enabled.status_code == 200
        assert client.get("/api/notifications/", headers=h).status_code == 200
        assert client.get("/api/auth/me", headers=h).json()["mfa_setup_required"] is False
    finally:
        client.delete("/api/admin/settings/security.require_mfa_staff", headers=admin)
        runtime_settings.invalidate_cache()


def test_captcha_sandbox_challenge_gates_login_and_registration():
    admin, _, _ = _auth_headers("admin")
    assert client.put("/api/admin/settings/security.captcha_enabled", json={"value": True},
                       headers=admin).status_code == 200
    try:
        challenge = client.get("/api/auth/captcha").json()
        assert challenge["provider"] == "sandbox" and "captcha_id" in challenge and "question" in challenge
        a, b = (int(x) for x in challenge["question"].replace("What is ", "").replace("?", "").split(" + "))

        # registration without solving the captcha is rejected
        payload = {"email": f"cap{uuid4().hex[:6]}@test.com", "password": PW, "full_name": "Cap Tester"}
        assert client.post("/api/auth/register", json=payload).status_code == 400
        # wrong answer is rejected
        bad = dict(payload, captcha_id=challenge["captcha_id"], captcha_answer=str(a + b + 1))
        assert client.post("/api/auth/register", json=bad).status_code == 400
        # correct answer succeeds
        good = dict(payload, captcha_id=challenge["captcha_id"], captcha_answer=str(a + b))
        assert client.post("/api/auth/register", json=good).status_code == 201

        # login is gated the same way
        email, _ = create_user("citizen")
        assert _login(email).status_code == 400
        challenge2 = client.get("/api/auth/captcha").json()
        a2, b2 = (int(x) for x in challenge2["question"].replace("What is ", "").replace("?", "").split(" + "))
        r = client.post("/api/auth/login", json={
            "email": email, "password": PW,
            "captcha_id": challenge2["captcha_id"], "captcha_answer": str(a2 + b2),
        })
        assert r.status_code == 200

        # a challenge can't be replayed with a different (or the same, twice-checked) id after tampering
        tampered = client.get("/api/auth/captcha").json()
        assert client.post("/api/auth/login", json={
            "email": email, "password": PW,
            "captcha_id": tampered["captcha_id"] + "x", "captcha_answer": "0",
        }).status_code == 400
    finally:
        client.delete("/api/admin/settings/security.captcha_enabled", headers=admin)
        runtime_settings.invalidate_cache()


def test_totp_matches_rfc6238_vector():
    import base64

    secret = base64.b32encode(b"12345678901234567890").decode()
    assert totp.totp_now(secret, at=59)[-6:] == "287082"  # RFC 6238 SHA-1 test vector (T=59)


def test_mfa_secret_is_encrypted_at_rest():
    h, email, uid = _auth_headers("citizen")
    secret = client.post("/api/auth/mfa/setup", headers=h).json()["secret"]
    db = _db()
    try:
        stored = db.query(User).filter(User.id == uid).first().mfa_secret_enc
    finally:
        db.close()
    assert stored and secret not in stored


def test_change_password_revokes_other_sessions():
    email, _ = create_user("citizen")
    t1 = _login(email).json()["access_token"]
    t2 = _login(email, **{"User-Agent": "other"}).json()["access_token"]
    h1, h2 = {"Authorization": f"Bearer {t1}"}, {"Authorization": f"Bearer {t2}"}
    bad = client.post("/api/auth/change-password", json={"current_password": "wrong", "new_password": "newpass123"}, headers=h1)
    assert bad.status_code == 400
    ok = client.post("/api/auth/change-password", json={"current_password": PW, "new_password": "newpass123"}, headers=h1)
    assert ok.status_code == 204
    assert client.get("/api/auth/me", headers=h2).status_code == 401
    assert client.get("/api/auth/me", headers=h1).status_code == 200
    assert _login(email, "newpass123").status_code == 200


def test_security_headers_and_request_id():
    r = client.get("/health")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["x-frame-options"] == "DENY"
    assert r.headers["x-request-id"] and "ms" in r.headers["x-process-time"]
    assert "content-security-policy" in client.get("/").headers


# ============================ RBAC / admin ============================

def test_citizen_cannot_reach_admin_endpoints():
    h, _, _ = _auth_headers("citizen")
    for path in ("/api/admin/users", "/api/admin/settings", "/api/admin/audit-logs", "/api/system/metrics",
                 "/api/reports/", "/api/integration/agencies"):
        assert client.get(path, headers=h).status_code == 403, path


def test_admin_creates_staff_and_manages_users():
    admin, _, admin_id = _auth_headers("admin")
    email = f"off{uuid4().hex[:6]}@rtsa.gov.zm"
    r = client.post("/api/admin/users", json={"email": email, "password": PW, "full_name": "New Officer",
                                              "role": "officer"}, headers=admin)
    assert r.status_code == 201 and r.json()["role"] == "officer"
    uid = r.json()["id"]
    assert _login(email).status_code == 200
    # deactivating signs them out and blocks login
    assert client.patch(f"/api/admin/users/{uid}", json={"is_active": False}, headers=admin).status_code == 200
    assert _login(email).status_code == 403
    # admins can't lock themselves out
    assert client.patch(f"/api/admin/users/{admin_id}", json={"is_active": False}, headers=admin).status_code == 400
    assert client.patch(f"/api/admin/users/{admin_id}/role?new_role=citizen", headers=admin).status_code == 400


def test_cannot_remove_the_last_admin():
    from app.core import permissions
    from app.models.user import User, UserRole

    admin, _, admin_id = _auth_headers("admin")
    other_admin, _, other_id = _auth_headers("admin")

    # The test DB is shared across the whole session, so other tests may have left
    # admin accounts lying around. Deactivate every admin except our two so the
    # "last admin" scenario below is deterministic regardless of run order.
    db = _db()
    try:
        others = (
            db.query(User)
            .filter(User.role == UserRole.ADMIN, User.is_active.is_(True), User.id.notin_([admin_id, other_id]))
            .all()
        )
        stashed_ids = [u.id for u in others]
        for u in others:
            u.is_active = False
        db.commit()
    finally:
        db.close()

    try:
        # two active admins exist, so a third party may demote/deactivate either one
        assert client.patch(f"/api/admin/users/{other_id}/role?new_role=officer", headers=admin).status_code == 200
        assert client.patch(f"/api/admin/users/{other_id}", json={"role": "admin"}, headers=admin).status_code == 200
        assert client.patch(f"/api/admin/users/{other_id}", json={"is_active": False}, headers=admin).status_code == 200
        assert client.patch(f"/api/admin/users/{other_id}", json={"is_active": True, "role": "admin"},
                             headers=admin).status_code == 200
        # now demote `other_id` for good, leaving `admin_id` as the sole active admin
        assert client.patch(f"/api/admin/users/{other_id}/role?new_role=officer", headers=admin).status_code == 200
        # a third party with elevated permissions (but not the admin role itself) still can't
        # demote or deactivate the last remaining admin, via either endpoint
        officer, _, _ = _auth_headers("officer")
        assert client.put("/api/admin/permissions/officer/users:manage?granted=true", headers=admin).status_code == 200
        assert client.put("/api/admin/permissions/officer/roles:manage?granted=true", headers=admin).status_code == 200
        permissions.invalidate_cache()
        try:
            assert client.patch(f"/api/admin/users/{admin_id}/role?new_role=officer",
                                 headers=officer).status_code == 400
            assert client.patch(f"/api/admin/users/{admin_id}", json={"is_active": False},
                                 headers=officer).status_code == 400
            assert client.patch(f"/api/admin/users/{admin_id}", json={"role": "officer"},
                                 headers=officer).status_code == 400
        finally:
            client.put("/api/admin/permissions/officer/users:manage?granted=false", headers=admin)
            client.put("/api/admin/permissions/officer/roles:manage?granted=false", headers=admin)
            permissions.invalidate_cache()
    finally:
        db = _db()
        try:
            for u in db.query(User).filter(User.id.in_(stashed_ids)).all():
                u.is_active = True
            db.commit()
        finally:
            db.close()


def test_permission_matrix_can_be_edited():
    admin, _, _ = _auth_headers("admin")
    officer, _, _ = _auth_headers("officer")
    assert client.get("/api/payments/summary", headers=officer).status_code == 403
    r = client.put("/api/admin/permissions/officer/payments:view_all?granted=true", headers=admin)
    assert r.status_code == 200 and "payments:view_all" in r.json()["permissions"]
    assert client.get("/api/payments/summary", headers=officer).status_code == 200
    client.put("/api/admin/permissions/officer/payments:view_all?granted=false", headers=admin)
    from app.core import permissions

    permissions.invalidate_cache()
    assert client.get("/api/payments/summary", headers=officer).status_code == 403
    assert client.put("/api/admin/permissions/admin/users:manage?granted=false", headers=admin).status_code == 400
    assert client.put("/api/admin/permissions/officer/nope:nope?granted=true", headers=admin).status_code == 404


def test_system_settings_validate_and_apply():
    admin, _, _ = _auth_headers("admin")
    settings_list = client.get("/api/admin/settings", headers=admin).json()
    assert any(s["key"] == "security.max_login_attempts" for s in settings_list)
    assert client.put("/api/admin/settings/security.max_login_attempts", json={"value": 1}, headers=admin).status_code == 422
    assert client.put("/api/admin/settings/nope", json={"value": 1}, headers=admin).status_code == 422
    r = client.put("/api/admin/settings/security.max_login_attempts", json={"value": 3}, headers=admin)
    assert r.status_code == 200
    # the new threshold takes effect immediately
    email, _ = create_user("citizen")
    for i in range(3):
        _login(email, "wrong-pass1", **{"X-Forwarded-For": f"10.1.0.{i}"})
    assert _login(email, PW, **{"X-Forwarded-For": "10.1.0.99"}).status_code == 423
    assert client.delete("/api/admin/settings/security.max_login_attempts", headers=admin).status_code == 204


def test_audit_log_records_admin_actions():
    admin, _, _ = _auth_headers("admin")
    client.post("/api/admin/users", json={"email": f"a{uuid4().hex[:6]}@t.com", "password": PW, "full_name": "Audit Me"},
                headers=admin)
    logs = client.get("/api/admin/audit-logs?action=create_user", headers=admin).json()
    assert logs and logs[0]["entity_type"] == "user"


# ============================ payments ============================

def test_fine_payment_uses_server_amount_and_issues_receipt():
    h, _, uid = _auth_headers("citizen")
    _rule("payment_receipt")
    _, challan_id = _vehicle_with_fine(uid, amount=75000)
    # a lower client-supplied amount is rejected
    r = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id, "amount": 1}, headers=h)
    assert r.status_code == 400
    r = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id}, headers=h)
    assert r.status_code == 201, r.text
    pay = r.json()
    assert pay["amount"] == 75000 and pay["status"] == "completed" and pay["receipt_number"].startswith("RCT-")
    db = _db()
    try:
        assert db.query(Challan).filter(Challan.id == challan_id).first().status == ChallanStatus.PAID
    finally:
        db.close()
    # paying twice is refused
    assert client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id},
                       headers=h).status_code == 409
    receipt = client.get(f"/api/payments/{pay['id']}/receipt", headers=h)
    assert receipt.status_code == 200 and receipt.json()["amount"] == 75000
    pdf = client.get(f"/api/payments/{pay['id']}/receipt.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    events = client.get(f"/api/payments/{pay['id']}/events", headers=h).json()
    assert [e["event_type"] for e in events] == ["initiated", "completed"]


def test_cannot_pay_someone_elses_fine_or_read_their_payment():
    owner, _, owner_id = _auth_headers("citizen")
    stranger, _, _ = _auth_headers("citizen")
    _, challan_id = _vehicle_with_fine(owner_id)
    assert client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id},
                       headers=stranger).status_code == 404
    pay = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id}, headers=owner).json()
    assert client.get(f"/api/payments/{pay['id']}", headers=stranger).status_code == 404
    assert client.get("/api/payments/?all=true", headers=stranger).status_code == 403
    mine = client.get("/api/payments/", headers=stranger).json()
    assert all(p["id"] != pay["id"] for p in mine)


def test_legacy_pay_challan_endpoint_now_goes_through_ledger():
    h, _, uid = _auth_headers("citizen")
    _, challan_id = _vehicle_with_fine(uid)
    assert client.post(f"/api/enforcement/challans/{challan_id}/pay", headers=h).status_code == 200
    payments = client.get("/api/payments/", headers=h).json()
    assert any(p["related_entity_id"] == challan_id for p in payments)
    stranger, _, _ = _auth_headers("citizen")
    _, other = _vehicle_with_fine(uid)
    assert client.post(f"/api/enforcement/challans/{other}/pay", headers=stranger).status_code == 404


def test_idempotency_key_prevents_double_charge():
    h, _, _ = _auth_headers("citizen")
    key = uuid4().hex
    body = {"payment_type": "fee", "amount": 12000, "description": "Vehicle transfer fee"}
    first = client.post("/api/payments/", json=body, headers={**h, "Idempotency-Key": key})
    second = client.post("/api/payments/", json=body, headers={**h, "Idempotency-Key": key})
    assert first.status_code == 201 and second.status_code == 200
    assert first.json()["id"] == second.json()["id"]
    other, _, _ = _auth_headers("citizen")
    assert client.post("/api/payments/", json=body, headers={**other, "Idempotency-Key": key}).status_code == 409


def test_fee_amount_validation():
    h, _, _ = _auth_headers("citizen")
    assert client.post("/api/payments/", json={"payment_type": "fee"}, headers=h).status_code == 400
    assert client.post("/api/payments/", json={"payment_type": "fee", "amount": -5}, headers=h).status_code == 422
    assert client.post("/api/payments/", json={"payment_type": "fee", "amount": 10**12}, headers=h).status_code == 400


def test_declined_payment_leaves_fine_unpaid():
    h, _, uid = _auth_headers("citizen")
    _, challan_id = _vehicle_with_fine(uid)
    r = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id,
                                            "gateway": "sandbox_decline"}, headers=h)
    assert r.status_code == 201 and r.json()["status"] == "failed" and r.json()["receipt_number"] is None
    db = _db()
    try:
        assert db.query(Challan).filter(Challan.id == challan_id).first().status == ChallanStatus.UNPAID
    finally:
        db.close()
    assert client.get(f"/api/payments/{r.json()['id']}/receipt", headers=h).status_code == 409
    # and it can be retried successfully
    assert client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id},
                       headers=h).json()["status"] == "completed"


def test_mobile_money_webhook_settles_asynchronously():
    h, _, uid = _auth_headers("citizen")
    _, challan_id = _vehicle_with_fine(uid)
    r = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id,
                                            "gateway": "mobile_money"}, headers=h)
    pay = r.json()
    assert pay["status"] == "pending"
    # no double-charging while a payment is in flight
    assert client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id,
                                               "gateway": "mobile_money"}, headers=h).status_code == 409
    body = json.dumps({"gateway_reference": pay["gateway_reference"], "outcome": "success"}).encode()
    unsigned = client.post("/api/payments/webhook/mobile_money", content=body)
    assert unsigned.status_code == 401
    forged = client.post("/api/payments/webhook/mobile_money", content=body, headers={"X-Signature": "deadbeef"})
    assert forged.status_code == 401
    from app.services.payments import sign_webhook

    ok = client.post("/api/payments/webhook/mobile_money", content=body,
                     headers={"X-Signature": sign_webhook("mobile_money", body)})
    assert ok.status_code == 200 and ok.json()["status"] == "completed"
    # replaying the callback is harmless
    again = client.post("/api/payments/webhook/mobile_money", content=body,
                        headers={"X-Signature": sign_webhook("mobile_money", body)})
    assert again.status_code == 200
    assert client.get(f"/api/payments/{pay['id']}/receipt", headers=h).status_code == 200


def test_refund_reopens_fine_and_is_permission_gated():
    admin, _, _ = _auth_headers("admin")
    h, _, uid = _auth_headers("citizen")
    _, challan_id = _vehicle_with_fine(uid, amount=40000)
    pay = client.post("/api/payments/", json={"payment_type": "fine", "related_entity_id": challan_id}, headers=h).json()
    assert client.post(f"/api/payments/{pay['id']}/refund", json={"reason": "mistake"}, headers=h).status_code == 403
    assert client.post(f"/api/payments/{pay['id']}/refund", json={"amount": 999999, "reason": "too much"},
                       headers=admin).status_code == 400
    part = client.post(f"/api/payments/{pay['id']}/refund", json={"amount": 10000, "reason": "partial"}, headers=admin)
    assert part.status_code == 200 and part.json()["status"] == "completed" and part.json()["refunded_amount"] == 10000
    full = client.post(f"/api/payments/{pay['id']}/refund", json={"reason": "rest"}, headers=admin)
    assert full.json()["status"] == "refunded" and full.json()["refunded_amount"] == 40000
    db = _db()
    try:
        assert db.query(Challan).filter(Challan.id == challan_id).first().status == ChallanStatus.UNPAID
    finally:
        db.close()


def test_reconciliation_finds_matches_and_discrepancies():
    admin, _, _ = _auth_headers("admin")
    h, _, _ = _auth_headers("citizen")
    refs = []
    for amount in (1000, 2000, 3000):
        p = client.post("/api/payments/", json={"payment_type": "fee", "amount": amount}, headers=h).json()
        refs.append((p["gateway_reference"], amount))
    statement = [
        {"gateway_reference": refs[0][0], "amount": refs[0][1]},           # matches
        {"gateway_reference": refs[1][0], "amount": refs[1][1] + 500},     # amount differs
        {"gateway_reference": "SANDBOX-UNKNOWN", "amount": 700},           # not in our ledger
    ]                                                                       # refs[2] missing from statement
    r = client.post("/api/payments/reconciliation/run", json={"gateway": "sandbox", "statement": statement}, headers=admin)
    assert r.status_code == 201, r.text
    run = r.json()
    assert run["matched"] >= 1 and run["mismatched"] >= 1
    assert run["missing_in_ledger"] == 1 and run["missing_in_statement"] >= 1
    assert any(m["gateway_reference"] == refs[1][0] for m in run["details"]["mismatched"])
    # a matched payment is not reconciled twice
    again = client.post("/api/payments/reconciliation/run",
                        json={"gateway": "sandbox", "statement": statement[:1]}, headers=admin).json()
    assert again["matched"] == 0
    assert client.get("/api/payments/reconciliation/runs", headers=admin).status_code == 200
    assert client.post("/api/payments/reconciliation/run", json={"gateway": "sandbox", "statement": []},
                       headers=h).status_code == 403


def test_revenue_summary():
    admin, _, _ = _auth_headers("admin")
    h, _, _ = _auth_headers("citizen")
    client.post("/api/payments/", json={"payment_type": "permit", "amount": 5000}, headers=h)
    s = client.get("/api/payments/summary?days=1", headers=admin).json()
    assert s["gross"] >= 5000 and s["net"] == s["gross"] - s["refunded"] and "permit" in s["by_type"]


# ============================ notifications ============================

def test_notification_preferences_and_delivery_pipeline():
    _rule("test_multi", "in_app,sms,email", "Hello", "Ref {reference}")
    h, _, uid = _auth_headers("citizen", phone_number="+260971234567")
    admin, _, _ = _auth_headers("admin")
    from app.services.notifications import notify

    db = _db()
    try:
        made = notify(db, uid, "test_multi", {"reference": "R1"})
        db.commit()
        assert {n.channel.value for n in made} == {"in_app", "sms", "email"}
        assert {n.status.value for n in made if n.channel.value != "in_app"} == {"pending"}
    finally:
        db.close()
    # in-app feed shows only the in-app row
    feed = client.get("/api/notifications/", headers=h).json()
    assert [n["channel"] for n in feed] == ["in_app"]
    dispatched = client.post("/api/system/dispatch-notifications", headers=admin).json()
    assert dispatched["sent"] >= 2  # sandbox providers accept
    # opt out of SMS
    prefs = client.put("/api/notifications/preferences", json={"sms": False}, headers=h).json()
    assert prefs["sms"] is False and prefs["email"] is True
    assert client.put("/api/notifications/preferences", json={"in_app": False}, headers=h).status_code == 422
    db = _db()
    try:
        made = notify(db, uid, "test_multi", {"reference": "R2"})
        db.commit()
        assert {n.channel.value for n in made} == {"in_app", "email"}
    finally:
        db.close()
    assert client.post("/api/notifications/read-all", headers=h).json()["marked"] >= 1
    assert client.get("/api/notifications/unread-count", headers=h).json()["unread"] == 0


def test_failed_delivery_is_retried_then_marked_failed(monkeypatch):
    _rule("test_retry", "email", "Hi", "There")
    _, _, uid = _auth_headers("citizen")
    from app.services import delivery
    from app.services.notifications import notify

    def boom(*a, **k):
        raise RuntimeError("smtp down")

    monkeypatch.setattr(delivery, "send_email", boom)
    db = _db()
    try:
        made = notify(db, uid, "test_retry", {})
        db.commit()
        nid = made[0].id
        for _ in range(delivery.MAX_ATTEMPTS):
            delivery.dispatch_pending(db)
        from app.models.notification import Notification

        n = db.query(Notification).filter(Notification.id == nid).first()
        assert n.status.value == "failed" and n.attempts == delivery.MAX_ATTEMPTS and "smtp down" in n.last_error
    finally:
        db.close()


def test_expiry_reminders_fire_once_per_threshold():
    from app.models.insurance import Insurance
    from app.services.expiry import scan_expiries

    _rule("insurance_expiring", "in_app", "Insurance", "{registration} expires {expiry_date} ({days_left}d)")
    _, _, uid = _auth_headers("citizen")
    vid, _ = _vehicle_with_fine(uid)
    db = _db()
    try:
        db.add(Insurance(vehicle_id=vid, provider="Madison", policy_number=f"P-{uuid4().hex[:8]}",
                         start_date=datetime.utcnow() - timedelta(days=300),
                         end_date=datetime.utcnow() + timedelta(days=6), is_active=True))
        db.commit()
        first = scan_expiries(db)
        second = scan_expiries(db)
        assert first["insurance"] >= 1 and second["insurance"] == 0
        # time passes: crossing into the 1-day bucket sends a fresh reminder
        third = scan_expiries(db, now=datetime.utcnow() + timedelta(days=5, hours=1))
        assert third["insurance"] >= 1
    finally:
        db.close()


def test_admin_broadcast_and_rule_management():
    admin, _, _ = _auth_headers("admin")
    citizen, _, _ = _auth_headers("citizen")
    r = client.post("/api/admin/notifications/broadcast", json={"title": "Planned outage", "message": "Sunday 02:00",
                                                                "roles": ["citizen"]}, headers=admin)
    assert r.status_code == 200 and r.json()["recipients"] >= 1
    assert any(n["title"] == "Planned outage" for n in client.get("/api/notifications/", headers=citizen).json())
    assert client.post("/api/admin/notifications/broadcast", json={"title": "x"}, headers=admin).status_code == 422
    assert client.get("/api/admin/notifications/stats", headers=admin).status_code == 200


# ============================ citizen portal ============================

def test_citizen_portal_views():
    h, email, uid = _auth_headers("citizen")
    vid, challan_id = _vehicle_with_fine(uid, amount=20000)
    dash = client.get("/api/citizen/dashboard", headers=h).json()
    assert len(dash["vehicles"]) == 1 and dash["fines"]["total_unpaid"] == 1 and dash["fines"]["total_amount"] == 20000
    summary = client.get("/api/citizen/summary", headers=h).json()
    assert summary["unpaid_fines"] == 1 and summary["amount_due"] == 20000
    detail = client.get(f"/api/citizen/vehicles/{vid}", headers=h).json()
    assert detail["compliant"] is False and any("insurance" in i.lower() for i in detail["issues"])
    assert client.get(f"/api/citizen/fines/{challan_id}", headers=h).json()["violation"]["type"] == "speeding"
    other, _, _ = _auth_headers("citizen")
    assert client.get(f"/api/citizen/vehicles/{vid}", headers=other).status_code == 404
    assert client.get(f"/api/citizen/fines/{challan_id}", headers=other).status_code == 404
    assert client.get("/api/citizen/dashboard", headers=other).json()["vehicles"] == []


def test_application_tracking_timeline():
    h, _, _ = _auth_headers("citizen")
    officer, _, _ = _auth_headers("officer")
    r = client.post("/api/licence-applications/", json={"first_name": "Ann", "last_name": "Phiri", "id_number": "111111/11/1",
                                                        "date_of_birth": "1995-05-05T00:00:00", "requested_class": "B"}, headers=h)
    assert r.status_code == 201, r.text
    apps = client.get("/api/citizen/applications", headers=h).json()
    assert len(apps) == 1 and apps[0]["status"] == "submitted"
    assert [s["state"] for s in apps[0]["steps"]][:2] == ["current", "pending"]
    assert apps[0]["next_step"] == "Theory test scheduled"
    detail = client.get(f"/api/citizen/applications/{apps[0]['id']}", headers=h)
    assert detail.status_code == 200
    other, _, _ = _auth_headers("citizen")
    assert client.get(f"/api/citizen/applications/{apps[0]['id']}", headers=other).status_code == 404


def test_profile_update_validates_phone():
    h, _, _ = _auth_headers("citizen")
    assert client.patch("/api/citizen/profile", json={"phone_number": "abc"}, headers=h).status_code == 422
    r = client.patch("/api/citizen/profile", json={"phone_number": "+260 971 234 567", "full_name": "New Name"}, headers=h)
    assert r.status_code == 200 and r.json()["full_name"] == "New Name"


# ============================ reports & exports ============================

def test_report_catalogue_dashboard_and_json():
    admin, _, _ = _auth_headers("admin")
    keys = {r["key"] for r in client.get("/api/reports/", headers=admin).json()}
    assert {"registrations", "licensing", "violations", "accidents", "psv", "revenue", "toll"} <= keys
    dash = client.get("/api/reports/dashboard", headers=admin).json()
    assert {"kpis", "trends", "breakdowns"} <= set(dash) and "violations" in dash["kpis"]
    for key in keys:
        r = client.get(f"/api/reports/{key}", headers=admin)
        assert r.status_code == 200, (key, r.text)
        assert r.json()["columns"] and isinstance(r.json()["summary"], dict)
    assert client.get("/api/reports/nope", headers=admin).status_code == 404


@pytest.mark.parametrize("fmt,magic", [("csv", b"\xef\xbb\xbf"), ("xlsx", b"PK"), ("pdf", b"%PDF")])
def test_report_exports(fmt, magic):
    admin, _, _ = _auth_headers("admin")
    r = client.get(f"/api/reports/violations?format={fmt}&date_from=2020-01-01", headers=admin)
    assert r.status_code == 200, r.text
    assert r.content.startswith(magic)
    assert "attachment" in r.headers["content-disposition"] and fmt in r.headers["content-disposition"]


def test_export_is_permission_gated_and_audited():
    admin, _, _ = _auth_headers("admin")
    toll_op, _, _ = _auth_headers("toll_operator")  # has reports:view but not reports:export
    assert client.get("/api/reports/toll", headers=toll_op).status_code == 200
    assert client.get("/api/reports/toll?format=csv", headers=toll_op).status_code == 403
    client.get("/api/reports/toll?format=csv", headers=admin)
    logs = client.get("/api/admin/audit-logs?action=export_report", headers=admin).json()
    assert logs and logs[0]["entity_id"] == "toll"


def test_csv_export_neutralises_formula_injection():
    from app.services import exporters

    out = exporters.to_csv(["a"], [["=HYPERLINK(\"http://evil\")"], ["+cmd"], ["normal"], [-5]]).decode("utf-8-sig")
    lines = out.strip().splitlines()
    assert lines[1].startswith("'=") or lines[1].startswith("\"'=")
    assert "'+cmd" in out and "normal" in out


def test_xlsx_export_is_readable():
    from openpyxl import load_workbook

    admin, _, _ = _auth_headers("admin")
    r = client.get("/api/reports/revenue?format=xlsx&date_from=2020-01-01", headers=admin)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0].startswith("Revenue") and "Summary" in wb.sheetnames


# ============================ inter-agency integration ============================

def _register_agency(admin, agency_type="insurance", **extra):
    payload = {"name": f"Agency {uuid4().hex[:6]}", "agency_type": agency_type, **extra}
    r = client.post("/api/integration/agencies", json=payload, headers=admin)
    assert r.status_code == 201, r.text
    return r.json()


def test_agency_key_shown_once_and_stored_hashed():
    admin, _, _ = _auth_headers("admin")
    agency = _register_agency(admin, "police")
    assert agency["api_key"].startswith("rtsa_")
    listed = client.get("/api/integration/agencies", headers=admin).json()
    row = next(a for a in listed if a["id"] == agency["id"])
    assert "api_key" not in row and agency["api_key"] not in json.dumps(row)


def test_agency_authentication_and_scope_enforcement():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0001")
    police = _register_agency(admin, "police", scopes=["vehicles:read"])
    hdr = {"X-API-Key": police["api_key"]}
    assert client.get("/api/integration/police/vehicles/INT 0001").status_code == 401
    assert client.get("/api/integration/police/vehicles/INT 0001", headers={"X-API-Key": "rtsa_bad_key"}).status_code == 401
    ok = client.get("/api/integration/police/vehicles/INT 0001", headers=hdr)
    assert ok.status_code == 200 and ok.json()["plate_number"] == "INT 0001"
    # contract doesn't include drivers:read or accidents:write
    assert client.get("/api/integration/police/drivers/X", headers=hdr).status_code == 403
    assert client.post("/api/integration/police/accidents", headers=hdr, json={
        "location": "x road", "occurred_at": "2026-01-01T10:00:00", "severity": "minor"}).status_code == 403
    # scopes can't exceed what the agency type may hold
    bad = client.post("/api/integration/agencies", headers=admin, json={
        "name": "Greedy", "agency_type": "hospital", "scopes": ["vehicles:read"]})
    assert bad.status_code == 422
    # JWT user tokens are not agency keys
    assert client.get("/api/integration/police/vehicles/INT 0001", headers=admin).status_code == 401


def test_disabled_expired_and_rotated_keys_stop_working():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0002")
    a = _register_agency(admin, "police")
    hdr = {"X-API-Key": a["api_key"]}
    url = "/api/integration/police/vehicles/INT 0002"
    assert client.get(url, headers=hdr).status_code == 200
    client.patch(f"/api/integration/agencies/{a['id']}", json={"is_active": False}, headers=admin)
    assert client.get(url, headers=hdr).status_code == 403
    client.patch(f"/api/integration/agencies/{a['id']}", json={"is_active": True}, headers=admin)
    assert client.get(url, headers=hdr).status_code == 200
    past = (datetime.utcnow() - timedelta(days=1)).isoformat()
    client.patch(f"/api/integration/agencies/{a['id']}", json={"contract_expires_at": past}, headers=admin)
    assert client.get(url, headers=hdr).status_code == 403
    client.patch(f"/api/integration/agencies/{a['id']}", json={"contract_expires_at": (datetime.utcnow() + timedelta(days=30)).isoformat()}, headers=admin)
    new_key = client.post(f"/api/integration/agencies/{a['id']}/rotate-key", headers=admin).json()["api_key"]
    assert client.get(url, headers=hdr).status_code == 401
    assert client.get(url, headers={"X-API-Key": new_key}).status_code == 200


def test_agency_rate_limit():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0003")
    a = _register_agency(admin, "police", rate_limit_per_minute=3)
    hdr = {"X-API-Key": a["api_key"]}
    codes = [client.get("/api/integration/police/vehicles/INT 0003", headers=hdr).status_code for _ in range(5)]
    assert codes == [200, 200, 200, 429, 429]


def test_insurer_pushes_policy_and_compliance_updates():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0004")
    ins = _register_agency(admin, "insurance")
    toll = _register_agency(admin, "toll_authority")
    hi, ht = {"X-API-Key": ins["api_key"]}, {"X-API-Key": toll["api_key"]}
    before = client.get("/api/integration/toll-authority/compliance/INT 0004", headers=ht).json()
    assert any("insurance" in i.lower() for i in before["issues"])
    policy = {"plate_number": "INT 0004", "provider": "ZSIC", "policy_number": f"Z-{uuid4().hex[:8]}",
              "start_date": (datetime.utcnow() - timedelta(days=1)).isoformat(),
              "end_date": (datetime.utcnow() + timedelta(days=200)).isoformat()}
    r = client.put("/api/integration/insurance/policies", json=policy, headers=hi)
    assert r.status_code == 200 and r.json()["created"] is True
    after = client.get("/api/integration/toll-authority/compliance/INT 0004", headers=ht).json()
    assert not any("insurance" in i.lower() for i in after["issues"])
    # idempotent update + validation
    assert client.put("/api/integration/insurance/policies", json=policy, headers=hi).json()["created"] is False
    bad = {**policy, "end_date": policy["start_date"]}
    assert client.put("/api/integration/insurance/policies", json=bad, headers=hi).status_code == 422
    assert client.put("/api/integration/insurance/policies", json={**policy, "plate_number": "NOPE 1"}, headers=hi).status_code == 404
    # the toll authority cannot write policies
    assert client.put("/api/integration/insurance/policies", json=policy, headers=ht).status_code == 403


def test_hospital_and_police_report_accidents():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0005")
    hosp = _register_agency(admin, "hospital")
    r = client.post("/api/integration/hospital/accidents", headers={"X-API-Key": hosp["api_key"]}, json={
        "location": "Kafue Rd", "occurred_at": "2026-09-01T08:30:00", "severity": "serious",
        "plate_numbers": ["int 0005", "GHOST 1"], "casualties": 2, "description": "Two admitted"})
    assert r.status_code == 201, r.text
    assert r.json()["linked_vehicles"] == ["INT 0005"] and r.json()["unrecognised_plates"] == ["GHOST 1"]
    accident = client.get(f"/api/accidents/{r.json()['accident_id']}", headers=admin)
    assert accident.status_code == 200 and "Reported by" in accident.json()["description"]


def test_integration_monitoring_records_calls_and_failures():
    admin, _, _ = _auth_headers("admin")
    _, uid = create_user("citizen")
    _vehicle_with_fine(uid, plate="INT 0006")
    a = _register_agency(admin, "police", scopes=["vehicles:read"])
    hdr = {"X-API-Key": a["api_key"]}
    client.get("/api/integration/police/vehicles/INT 0006", headers=hdr)
    client.get("/api/integration/police/vehicles/UNKNOWN", headers=hdr)           # 404
    client.get("/api/integration/police/drivers/X", headers=hdr)                   # 403
    mon = client.get("/api/integration/monitoring", headers=admin).json()
    row = next(x for x in mon["agencies"] if x["agency"] == a["name"])
    assert row["calls"] == 3 and row["successes"] == 1 and row["failures"] == 2
    assert row["success_rate"] == pytest.approx(33.3, abs=0.1)
    assert any(e["agency"] == a["name"] for e in mon["recent_errors"])
    logs = client.get(f"/api/integration/logs?agency={a['name']}", headers=admin).json()
    assert len(logs) == 3 and {l["status"] for l in logs} == {200, 404, 403}


def test_national_id_verification_is_honest_in_sandbox():
    officer, _, _ = _auth_headers("officer")
    citizen, _, _ = _auth_headers("citizen")
    r = client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer)
    assert r.status_code == 200 and r.json()["verified"] is False and r.json()["source"] == "sandbox"
    assert client.get("/api/integration/national-id/verify?nrc=garbage", headers=officer).status_code == 422
    assert client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=citizen).status_code == 403


def test_national_id_real_provider_success_retry_and_bad_config(monkeypatch):
    from app.core.config import settings as env
    from app.services import integration as integration_service

    officer, _, _ = _auth_headers("officer")
    monkeypatch.setattr(env, "NATIONAL_ID_API_URL", "https://registry.example.gov.zm/verify")
    monkeypatch.setattr(env, "NATIONAL_ID_API_TOKEN", "test-token")
    monkeypatch.setattr(integration_service.time, "sleep", lambda *_: None)  # don't slow the test down

    class FakeResponse:
        def __init__(self, status_code, body):
            self.status_code = status_code
            self._body = body

        def json(self):
            if self._body is None:
                raise ValueError("no body")
            return self._body

    # happy path
    calls = []

    def ok(url, params, timeout, headers):
        calls.append(params["nrc"])
        assert headers["Authorization"] == "Bearer test-token"
        return FakeResponse(200, {"verified": True, "full_name": "Jane Banda"})

    monkeypatch.setattr(integration_service.httpx, "get", ok)
    r = client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer)
    assert r.status_code == 200
    body = r.json()
    assert body == {"nrc": "123456/78/1", "format_valid": True, "verified": True,
                     "source": "national_id_registry", "full_name": "Jane Banda"}

    # a transient 503 is retried and then succeeds
    attempts = {"n": 0}

    def flaky(url, params, timeout, headers):
        attempts["n"] += 1
        if attempts["n"] < 2:
            return FakeResponse(503, None)
        return FakeResponse(200, {"verified": False})

    monkeypatch.setattr(integration_service.httpx, "get", flaky)
    r = client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer)
    assert r.status_code == 200 and r.json()["verified"] is False and attempts["n"] == 2

    # persistent failure surfaces as a clean 502, not a crash
    def always_down(url, params, timeout, headers):
        return FakeResponse(503, None)

    monkeypatch.setattr(integration_service.httpx, "get", always_down)
    assert client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer).status_code == 502

    # a malformed (non-JSON / non-dict) response is also a clean 502
    def bad_body(url, params, timeout, headers):
        return FakeResponse(200, None)

    monkeypatch.setattr(integration_service.httpx, "get", bad_body)
    assert client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer).status_code == 502

    # missing token is caught as a misconfiguration, not attempted
    monkeypatch.setattr(env, "NATIONAL_ID_API_TOKEN", "")

    def should_not_be_called(*a, **k):
        raise AssertionError("must not call out with a missing token")

    monkeypatch.setattr(integration_service.httpx, "get", should_not_be_called)
    r = client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer)
    assert r.status_code == 500 and "TOKEN" in r.json()["detail"]

    # a non-https URL is refused in production
    monkeypatch.setattr(env, "NATIONAL_ID_API_TOKEN", "test-token")
    monkeypatch.setattr(env, "NATIONAL_ID_API_URL", "http://registry.example.gov.zm/verify")
    monkeypatch.setattr(env, "ENVIRONMENT", "production")
    r = client.get("/api/integration/national-id/verify?nrc=123456/78/1", headers=officer)
    assert r.status_code == 500 and "https" in r.json()["detail"]


# ============================ operations ============================

def test_health_probes():
    assert client.get("/health/live").json()["status"] == "alive"
    r = client.get("/health/ready")
    assert r.status_code == 200 and r.json()["status"] == "ready"


def test_metrics_endpoint_reports_latency_and_toll_target():
    admin, _, _ = _auth_headers("admin")
    for _ in range(3):
        client.get("/health")
    m = client.get("/api/system/metrics", headers=admin).json()
    assert m["total_requests"] > 0 and "GET /health" in m["routes"]
    assert m["routes"]["GET /health"]["p95_ms"] >= 0
    assert m["toll_compliance_target"]["target_ms"] == 500


def test_backup_create_verify_and_restore_roundtrip(tmp_path, monkeypatch):
    from app.core.config import settings
    from scripts import backup, restore

    monkeypatch.setattr(settings, "BACKUP_DIR", str(tmp_path))
    admin, _, _ = _auth_headers("admin")
    r = client.post("/api/system/backups", headers=admin)
    assert r.status_code == 200, r.text
    assert r.json()["verification"]["ok"] is True
    listing = client.get("/api/system/backups", headers=admin).json()
    assert listing["count"] == 1
    path = tmp_path / listing["latest"]["file"]

    with gzip.open(path, "rt") as fh:
        snapshot = json.load(fh)
    user_count = snapshot["counts"]["users"]
    assert user_count > 0

    # a truncated/corrupted file fails verification
    bad = tmp_path / "rtsa-itms-corrupt.json.gz"
    bad.write_bytes(path.read_bytes()[:40])
    assert backup.verify_backup(bad)["ok"] is False

    # restore refuses to clobber a non-empty database without --wipe, then round-trips with it
    with pytest.raises(SystemExit):
        restore.restore_json(path, wipe=False)
    counts = restore.restore_json(path, wipe=True)
    assert counts["users"] == user_count
    db = _db()
    try:
        assert db.query(User).count() == user_count
    finally:
        db.close()


def test_backup_retention_prunes_old_files(tmp_path):
    from scripts.backup import prune

    for i in range(5):
        f = tmp_path / f"rtsa-itms-{i}.json.gz"
        f.write_bytes(b"x")
        import os
        os.utime(f, (1000 + i, 1000 + i))
    assert len(prune(tmp_path, keep=2)) == 3
    assert len(list(tmp_path.iterdir())) == 2


def test_capacity_endpoint():
    admin, _, _ = _auth_headers("admin")
    r = client.get("/api/system/capacity", headers=admin).json()
    assert r["row_counts"]["vehicles"] >= 0 and "audit_logs" in r["row_counts"]


# ============================ hardening regressions ============================

def test_forwarded_for_is_ignored_unless_proxy_is_trusted(monkeypatch):
    from app.core.config import settings
    from starlette.requests import Request

    from app.services.auth import client_ip

    def req(xff):
        return Request({"type": "http", "headers": [(b"x-forwarded-for", xff.encode())], "client": ("9.9.9.9", 1)})

    monkeypatch.setattr(settings, "TRUST_PROXY_HEADERS", False)
    assert client_ip(req("1.1.1.1")) == "9.9.9.9"              # header ignored: cannot be spoofed
    monkeypatch.setattr(settings, "TRUST_PROXY_HEADERS", True)
    assert client_ip(req("6.6.6.6, 2.2.2.2")) == "2.2.2.2"       # rightmost = added by our proxy, not the client


def test_successful_login_does_not_reset_ip_failure_window():
    own, _ = create_user("citizen")
    victim, _ = create_user("citizen")
    hdr = {"X-Forwarded-For": "7.7.7.7"}
    for _ in range(3):
        _login(victim, "bad-guess-1", **hdr)
    assert _login(own, **hdr).status_code == 200            # attacker logs into their own account...
    for _ in range(2):
        _login(victim, "bad-guess-2", **hdr)
    assert _login(victim, "bad-guess-3", **hdr).status_code == 429  # ...but the IP is still throttled


def test_admin_password_reset_takes_body_not_query():
    admin, _, _ = _auth_headers("admin")
    email, uid = create_user("citizen")
    assert client.post(f"/api/admin/users/{uid}/reset-password", headers=admin).status_code == 422
    assert client.post(f"/api/admin/users/{uid}/reset-password", json={"new_password": "weak"}, headers=admin).status_code == 400
    assert client.post(f"/api/admin/users/{uid}/reset-password", json={"new_password": "brandnew123"}, headers=admin).status_code == 204
    assert _login(email, "brandnew123").status_code == 200


def test_malformed_ids_return_404_not_500():
    quiet = TestClient(app, raise_server_exceptions=False)
    email, _ = create_user("admin")
    token = quiet.post("/api/auth/login", json={"email": email, "password": PW}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    for path in ("/api/payments/not-a-uuid", "/api/vehicles/abc", "/api/citizen/vehicles/zzz", "/api/admin/users/xyz/sessions"):
        assert quiet.get(path, headers=h).status_code in (404, 200), path
    assert quiet.get("/api/payments/not-a-uuid", headers=h).status_code == 404


def test_maintenance_prunes_only_old_telemetry():
    from app.models.platform import LoginAttempt, UserSession
    from app.services.maintenance import prune_old_records

    email, uid = create_user("citizen")
    db = _db()
    try:
        old = datetime.utcnow() - timedelta(days=200)
        db.add(LoginAttempt(email=email, ip_address="1.1.1.1", success=False, created_at=old))
        db.add(LoginAttempt(email=email, ip_address="1.1.1.1", success=True))
        db.add(UserSession(user_id=uid, created_at=old, last_seen_at=old, expires_at=old, revoked_at=old))
        db.add(UserSession(user_id=uid, expires_at=datetime.utcnow() + timedelta(hours=1)))  # live: keep
        db.commit()
        result = prune_old_records(db)
        assert result["login_attempts"] >= 1 and result["sessions"] >= 1
        assert db.query(LoginAttempt).filter(LoginAttempt.email == email, LoginAttempt.success == True).count() == 1  # noqa: E712
        assert db.query(UserSession).filter(UserSession.user_id == uid, UserSession.revoked_at.is_(None)).count() == 1
    finally:
        db.close()


# ============================ toll-gate lookup ============================

def test_gate_lookup_lists_pending_offences_without_side_effects():
    from app.models.toll import TollTransaction

    officer, _, _ = _auth_headers("officer")
    _, uid = create_user("citizen")
    plate = f"LK {uuid4().hex[:4].upper()}"
    _vehicle_with_fine(uid, amount=60000, plate=plate)
    db = _db()
    before = (db.query(Challan).count(), db.query(TollTransaction).count())
    db.close()

    r = client.get(f"/api/toll/lookup/{plate.lower().replace(' ', '')}", headers=officer)   # spacing/case tolerant
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["found"] and body["decision"] == "stop"
    assert body["offence_count"] == 1 and body["total_owed"] == 60000
    assert body["offences"][0]["offence"] == "speeding" and body["offences"][0]["location"] == "Great East Rd"
    # looking someone up must not create fines or toll records - and can be repeated safely
    client.get(f"/api/toll/lookup/{plate}", headers=officer)
    db = _db()
    assert (db.query(Challan).count(), db.query(TollTransaction).count()) == before
    db.close()

    # once the fine is paid the offence disappears from the lookup
    client.post(f"/api/enforcement/challans/{body['offences'][0]['challan_id']}/pay", headers=officer)
    again = client.get(f"/api/toll/lookup/{plate}", headers=officer).json()
    assert again["offence_count"] == 0 and again["total_owed"] == 0


def test_gate_lookup_unknown_plate_and_permissions():
    officer, _, _ = _auth_headers("officer")
    citizen, _, _ = _auth_headers("citizen")
    r = client.get("/api/toll/lookup/NOPE999", headers=officer)
    assert r.status_code == 200 and r.json()["found"] is False and r.json()["decision"] == "stop"
    assert client.get("/api/toll/lookup/NOPE999", headers=citizen).status_code == 403
    assert client.get("/api/toll/lookup/NOPE999").status_code == 401
    assert client.get("/api/toll/lookup/ab", headers=officer).status_code == 422
    admin, _, _ = _auth_headers("admin")
    logs = client.get("/api/admin/audit-logs?action=plate_lookup", headers=admin).json()
    assert logs and "NOPE999" in logs[0]["details"]
