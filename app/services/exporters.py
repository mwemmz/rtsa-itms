"""Report and receipt exporters: CSV, Excel (.xlsx) and PDF."""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value) -> str:
    """Neutralise spreadsheet formula injection in exported text cells."""
    text = "" if value is None else str(value)
    if text.startswith(_FORMULA_PREFIXES) and not _is_number(text):
        return "'" + text
    return text


def _is_number(text: str) -> bool:
    try:
        float(text)
        return True
    except ValueError:
        return False


def _cell(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return value


def to_csv(columns: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_safe_text(_cell(v)) if isinstance(_cell(v), str) or v is None else _cell(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(title: str, columns: list[str], rows: list[list], summary: dict | None = None) -> bytes:
    # write_only streams rows to the file instead of holding the whole sheet in memory,
    # which keeps memory flat on large exports (measured: 50k rows ~8-10s, dominated by cell writing).
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31])
    for i, col in enumerate(columns, 1):
        width = max([len(str(col))] + [len(str(_cell(r[i - 1]) if i - 1 < len(r) else "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    header = []
    for col in columns:
        c = WriteOnlyCell(ws, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        header.append(c)
    ws.append(header)
    for row in rows:
        ws.append([_safe_text(_cell(v)) if isinstance(_cell(v), str) else _cell(v) for v in row])
    if summary:
        s = wb.create_sheet("Summary")
        s.column_dimensions["A"].width = 32
        s.column_dimensions["B"].width = 24
        for k, v in summary.items():
            key = WriteOnlyCell(s, value=str(k))
            key.font = Font(bold=True)
            s.append([key, _safe_text(v) if isinstance(v, str) else v])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(title: str, columns: list[str], rows: list[list], summary: dict | None = None,
           subtitle: str | None = None) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small")
    small.fontSize = 8
    small.leading = 10
    story = [Paragraph("RTSA Integrated Transport Management System", styles["Title"]),
             Paragraph(title, styles["Heading2"])]
    meta = f"Generated {datetime.utcnow():%Y-%m-%d %H:%M} UTC"
    if subtitle:
        meta += f" &middot; {subtitle}"
    story += [Paragraph(meta, small), Spacer(1, 6)]
    if summary:
        srows = [[Paragraph(f"<b>{_esc(k)}</b>", small), Paragraph(_esc(v), small)] for k, v in summary.items()]
        st = Table(srows, colWidths=[70 * mm, 60 * mm], hAlign="LEFT")
        st.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke)]))
        story += [st, Spacer(1, 8)]
    data = [[Paragraph(f"<b>{_esc(c)}</b>", small) for c in columns]]
    for row in rows:
        data.append([Paragraph(_esc(_cell(v)), small) for v in row])
    if not rows:
        data.append([Paragraph("No records for the selected period.", small)] + [""] * (len(columns) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")]),
    ]))
    story.append(table)
    doc.build(story)
    return out.getvalue()


def receipt_pdf(receipt: dict) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm, title=f"Receipt {receipt['receipt_number']}")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Road Transport &amp; Safety Agency", styles["Title"]),
        Paragraph("Official Payment Receipt", styles["Heading2"]),
        Spacer(1, 8),
    ]
    rows = [
        ["Receipt number", receipt["receipt_number"]],
        ["Payment reference", receipt["reference"]],
        ["Description", receipt.get("description") or receipt["payment_type"]],
        ["Payment type", receipt["payment_type"]],
        ["Amount", f"{receipt['currency']} {receipt['amount']:,}"],
        ["Refunded", f"{receipt['currency']} {receipt['refunded_amount']:,}"],
        ["Gateway", f"{receipt['gateway']} ({receipt.get('gateway_reference') or '-'})"],
        ["Paid by", receipt.get("paid_by_name") or "-"],
        ["Paid at", str(receipt.get("paid_at") or "-")],
        ["Status", receipt["status"]],
    ]
    table = Table([[Paragraph(f"<b>{_esc(a)}</b>", styles["BodyText"]), Paragraph(_esc(b), styles["BodyText"])]
                   for a, b in rows], colWidths=[50 * mm, 110 * mm])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                               ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke)]))
    story += [table, Spacer(1, 12),
              Paragraph("This receipt was generated electronically and is valid without a signature.",
                        styles["Italic"])]
    doc.build(story)
    return out.getvalue()


def _esc(value) -> str:
    return (str(value if value is not None else "")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
